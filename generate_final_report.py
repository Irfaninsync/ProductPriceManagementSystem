import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from report_builder_lib import (
    set_cell_background, set_cell_margins, set_table_borders, format_row,
    add_callout, add_heading_1, add_heading_2, add_heading_3, add_p, add_bullet
)

def build_docx():
    doc = Document()
    
    # Configure 1-inch margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
        
        # Header & Footer setup
        s.different_first_page_header_footer = True
        
        # Header (Pages 2+)
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Product Price Management System (PPMS) — Software Engineering Report")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(110, 118, 129)
        
        # Footer (Pages 2+)
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Core Java Enterprise Application  |  Page ")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(110, 118, 129)
        fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
        fp._element.append(fldSimple)

    # =========================================================================
    # COVER PAGE
    # =========================================================================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(24)
    p_inst.paragraph_format.space_after = Pt(8)
    r_inst = p_inst.add_run("ACADEMIC SOFTWARE ENGINEERING PROJECT REPORT\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(110, 118, 129)
    
    p_bar = doc.add_paragraph()
    p_bar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bar.paragraph_format.space_before = Pt(8)
    p_bar.paragraph_format.space_after = Pt(28)
    r_bar = p_bar.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    r_bar.font.name = "Calibri"
    r_bar.font.size = Pt(12)
    r_bar.font.color.rgb = RGBColor(43, 108, 176) # Slate Blue
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(10)
    r_title = p_title.add_run("PRODUCT PRICE MANAGEMENT SYSTEM\n(PPMS)")
    r_title.font.name = "Calibri Light"
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93) # Deep Navy
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(4)
    p_sub.paragraph_format.space_after = Pt(36)
    r_sub = p_sub.add_run("A Modular, Offline-First Desktop Management Engine Featuring Comprehensive Pricing Lifecycle Governance, Immutable Audit Trails, Analytical Valuation Intelligence, and Crash-Resilient CSV File Persistence in Core Java")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(74, 85, 104)
    
    # Team Registration Table on Cover Page
    p_team_hdr = doc.add_paragraph()
    p_team_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_team_hdr.paragraph_format.space_before = Pt(20)
    p_team_hdr.paragraph_format.space_after = Pt(8)
    r_th = p_team_hdr.add_run("SUBMITTED BY: ACADEMIC ENGINEERING PROJECT TEAM")
    r_th.font.name = "Calibri"
    r_th.font.size = Pt(10)
    r_th.font.bold = True
    r_th.font.color.rgb = RGBColor(27, 54, 93)
    
    tbl_cov = doc.add_table(rows=7, cols=2)
    tbl_cov.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cov.autofit = False
    tbl_cov.columns[0].width = Inches(3.0)
    tbl_cov.columns[1].width = Inches(3.0)
    set_table_borders(tbl_cov, color="CBD5E1", sz="4")
    
    cov_headers = ["Team Member Identifier", "Candidate Registration Number"]
    for i, h in enumerate(cov_headers):
        tbl_cov.cell(0, i).paragraphs[0].text = h
    format_row(tbl_cov.rows[0], "F1F5F9", is_header=True)
    
    cov_members = [
        ("Team Member 1 (Project Lead & Architect)", "2025DCSBT10650"),
        ("Team Member 2 (Product Model & Validation)", "2025DCSBT10651"),
        ("Team Member 3 (Product Operations & Search)", "2025DCSBT10652"),
        ("Team Member 4 (Pricing & Historical Audit)", "2025DCSBT10653"),
        ("Team Member 5 (Persistence & Quality Assurance)", "2025DCSBT10654"),
        ("Team Member 6 (Documentation & Demonstration)", "2025DCSBT10655"),
    ]
    for r_idx, (m_id, reg_no) in enumerate(cov_members):
        row = tbl_cov.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = m_id
        row.cells[1].paragraphs[0].text = reg_no
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")
        
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(40)
    p_meta.paragraph_format.space_after = Pt(0)
    r_meta = p_meta.add_run("Execution Environment: Standard Java SE 17 | Local File System Persistence\nAcademic Session 2025–2026")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_page_break()

    # =========================================================================
    # SECTION 1: EXECUTIVE ABSTRACT
    # =========================================================================
    add_heading_1(doc, "1. Executive Abstract")
    add_p(doc, 
        "In contemporary enterprise inventory administration and retail operations, conventional catalog software systems suffer from a severe architectural limitation: they maintain only the current state of merchandise prices while completely discarding the historical context of pricing revisions. When selling prices fluctuate due to wholesale supplier revaluations, inflationary pressures, promotional discounts, clearance sales, or competitor price matching, store managers and commercial administrators are left without historical traceability. Consequently, organizations struggle to identify when prices were adjusted, the commercial justification behind changes, and the direct impact on profitability margins.",
        bold_prefix="Context and Problem Scope: "
    )
    add_p(doc,
        "The Product Price Management System (PPMS) addresses this critical industry gap. PPMS is an advanced, standalone, offline-first Core Java software application architected to enforce strict pricing governance, maintain an immutable historical audit trail of all approved price modifications, generate comprehensive analytical valuation intelligence, and provide crash-safe persistence using local CSV files without requiring external database engines, network connectivity, or third-party enterprise frameworks.",
        bold_prefix="Engineered Solution: "
    )
    add_p(doc,
        "Architecturally, the application strictly adheres to a decoupled five-layer Object-Oriented design comprising Presentation, Service, Domain, Repository, and Persistence layers. Persisted storage relies upon custom RFC-4180 compliant finite-state machine CSV parsers and writers featuring atomic file replacement semantics (.tmp staging, .bak generation, and atomic renaming) to eliminate risks of data corruption. All financial computations strictly leverage java.math.BigDecimal with explicit RoundingMode.HALF_UP policies to eliminate binary floating-point roundoff errors.",
        bold_prefix="Methodology & Precision: "
    )
    add_p(doc,
        "The software engineering lifecycle followed a disciplined five-activity framework: Communication, Planning, Modelling, Construction, and Deployment. Quality assurance was verified through a self-contained test execution suite comprising 25 test cases across 6 distinct suites, achieving a verified 100% test pass rate covering domain models, pricing workflows, RFC-4180 parsing edge cases, financial analytics, batch CSV imports, and complete end-to-end restart persistence.",
        bold_prefix="Verification & Results: "
    )
    
    add_callout(doc, "Key Academic & Technical Contribution",
        "PPMS demonstrates that enterprise-grade software qualities—including ACID-like safe write operations, comprehensive audit trails, deterministic sorting, and high-watermark collision-free ID generation—can be successfully realized using pure Core Java standard libraries without relying on heavyweight frameworks or external DBMS engines.")

    # =========================================================================
    # SECTION 2: INTRODUCTION & BACKGROUND
    # =========================================================================
    add_heading_1(doc, "2. Introduction & Background")
    add_p(doc, "Product pricing represents one of the most critical operational levers governing commercial survival, profit margins, and inventory turnover in modern commerce. Whether operating a specialized electronics outlet, a regional wholesale distributor, or an academic retail cooperative, accurate product information and disciplined price management are paramount.")
    add_p(doc, "Traditional small-to-medium enterprise (SME) inventory management is overwhelmingly conducted through ad-hoc mechanisms such as manual paper ledgers, fragmented desktop spreadsheets, or rudimentary single-table catalog programs. While such tools permit basic record keeping, they fail entirely to provide pricing governance. In a typical spreadsheet, updating a price is executed by overwriting an existing cell. Once overwritten, the previous price, the date of alteration, and the business rationale for the modification are permanently lost.")
    add_p(doc, "The Product Price Management System (PPMS) is engineered specifically to eliminate these vulnerabilities. Operating entirely offline in a local Java runtime environment, PPMS provides an advanced, keyboard-driven console interface designed with a minimalist aesthetic, high-contrast readability, clear tabular formatting, and structured input validation. By separating product catalog operations from formal pricing decision workflows, PPMS ensures that every pricing modification is validated, previewed, justified, confirmed, and permanently recorded in an immutable audit ledger.")

    # =========================================================================
    # SECTION 3: PROBLEM STATEMENT
    # =========================================================================
    add_heading_1(doc, "3. Problem Statement & Academic Motivation")
    add_p(doc, "A critical examination of conventional desktop inventory recording tools identifies several persistent vulnerabilities that undermine commercial integrity:")
    add_bullet(doc, "Overwriting selling prices without historical logging erases critical trend data, rendering retrospective price elasticity audits impossible.", bold_prefix="Loss of Historical Pricing Context: ")
    add_bullet(doc, "In ad-hoc systems, prices can be modified arbitrarily without requiring justification, leaving management unable to verify whether updates were caused by supplier cost surges or unauthorized staff tampering.", bold_prefix="Absence of Operational Accountability: ")
    add_bullet(doc, "Standard desktop applications frequently store monetary values using primitive floating-point types (float, double), introducing binary truncation errors (e.g., 0.1 + 0.2 = 0.30000000000000004) that distort accounting reports.", bold_prefix="Financial Representation Inaccuracy: ")
    add_bullet(doc, "Directly overwriting persistent text files risks catastrophic data loss if power failure, operating system crash, or thread interruption occurs mid-stream. Furthermore, naive string splitting fails when records contain commas, quotes, or newlines.", bold_prefix="Data Corruption in Ad-Hoc Persistence: ")
    add_bullet(doc, "Store managers lack real-time visibility into products inadvertently sold below procurement cost (loss-leaders) or categories exhibiting gross margin decay.", bold_prefix="Deficit of Proactive Diagnostic Intelligence: ")
    add_p(doc, "Problem Statement Formalization: Commercial product administration requires a localized, autonomous software solution capable of strictly enforcing pricing governance, preventing binary calculation errors, ensuring collision-free record persistence, and providing complete historical auditability using standard Core Java technologies.")

    # =========================================================================
    # SECTION 4: PROJECT OBJECTIVES
    # =========================================================================
    add_heading_1(doc, "4. Project Objectives")
    add_p(doc, "The project objectives are partitioned into primary architectural imperatives and secondary advanced analytical enhancements:")
    add_heading_2(doc, "4.1 Primary Objectives")
    add_bullet(doc, "Implement a fully encapsulated Product catalog supporting creation, inspection, editing, multi-attribute searching, deterministic sorting, and soft archival.", bold_prefix="PO-1 (Catalog Governance): ")
    add_bullet(doc, "Develop a dedicated pricing workflow that computes signed difference and percentage change, mandates a business justification, requires explicit user confirmation, and persists an immutable audit record.", bold_prefix="PO-2 (Pricing Lifecycle & Audit): ")
    add_bullet(doc, "Design and construct an RFC-4180 compliant CSV parsing and serialization engine in pure Java with atomic file replacement mechanics to guarantee persistent integrity.", bold_prefix="PO-3 (Crash-Safe Persistence): ")
    add_bullet(doc, "Standardize all currency storage and financial calculations using java.math.BigDecimal with explicit half-up rounding, eliminating floating-point errors.", bold_prefix="PO-4 (Financial Precision): ")
    add_bullet(doc, "Build an interactive, highly readable console user interface featuring clean tabular alignment, pagination, and toggleable ANSI colors.", bold_prefix="PO-5 (Terminal UI/UX): ")
    
    add_heading_2(doc, "4.2 Secondary Objectives")
    add_bullet(doc, "Calculate real-time operational statistics including median selling price, inventory cost valuation, retail valuation, gross margin percentage, and markup percentage.", bold_prefix="SO-1 (Executive Analytics): ")
    add_bullet(doc, "Identify and flag inventory anomalies including products priced below cost and stock levels falling below a configurable threshold.", bold_prefix="SO-2 (Diagnostic Warning Alerts): ")
    add_bullet(doc, "Provide structured CSV export capabilities and robust batch CSV import with row-by-row pre-validation and rejection reporting.", bold_prefix="SO-3 (Batch Interoperability): ")
    add_bullet(doc, "Enable automated timestamped snapshot backup creation and pre-validated restoration with safety rollback mechanisms.", bold_prefix="SO-4 (Data Recovery): ")

    # =========================================================================
    # SECTION 5: PROJECT SCOPE
    # =========================================================================
    add_heading_1(doc, "5. Project Scope & Boundaries")
    add_p(doc, "Defining precise boundary conditions is essential to maintaining software engineering rigor and preventing scope creep:")
    
    add_heading_2(doc, "5.1 In-Scope Capabilities")
    add_bullet(doc, "Stand-alone Core Java SE 17 desktop execution on standard operating systems (Windows, Linux, macOS).")
    add_bullet(doc, "Interactive command-line interface featuring main dashboard, structured menus, paginated listings, and ASCII/box tables.")
    add_bullet(doc, "Complete product lifecycle management with permanent, collision-free alphanumeric IDs (PRD-xxxx).")
    add_bullet(doc, "Dedicated single and bulk pricing adjustment workflows with preview and user confirmation.")
    add_bullet(doc, "Chronological price history audit log with product, date-range, and keyword filtering.")
    add_bullet(doc, "Real-time calculation of financial valuation and health metrics (cost, retail, gross margin, markup, median).")
    add_bullet(doc, "Local persistence using RFC-4180 compliant CSV files with atomic write staging (.tmp) and backup (.bak).")
    add_bullet(doc, "Full automated verification test suite executed without external testing frameworks.")
    
    add_heading_2(doc, "5.2 Out-of-Scope Elements (Strict Technology Restrictions)")
    add_bullet(doc, "Web or mobile application frontends (HTML5, React, Angular, Vue, Flutter, Android).")
    add_bullet(doc, "Enterprise web frameworks and backends (Spring Boot, Jakarta EE, Node.js, Django).")
    add_bullet(doc, "External relational or NoSQL database servers (MySQL, PostgreSQL, Oracle, SQLite, MongoDB).")
    add_bullet(doc, "Object-Relational Mapping (ORM) tools such as Hibernate or JPA.")
    add_bullet(doc, "Cloud-hosted database services, REST APIs, or online payment gateways.")
    add_bullet(doc, "Graphical desktop GUI toolkits (Swing, JavaFX, AWT) as substitutes for the required console UI.")

    # =========================================================================
    # SECTION 6: EXISTING SYSTEM VS. PROPOSED PPMS
    # =========================================================================
    add_heading_1(doc, "6. Existing System vs. Proposed PPMS")
    add_p(doc, "To appreciate the technological and operational advancements delivered by PPMS, Table 6.1 contrasts the characteristics of traditional ad-hoc inventory methods against the engineered capabilities of the Product Price Management System.")
    
    tbl_comp = doc.add_table(rows=10, cols=3)
    tbl_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_comp.autofit = False
    tbl_comp.columns[0].width = Inches(1.8)
    tbl_comp.columns[1].width = Inches(2.3)
    tbl_comp.columns[2].width = Inches(2.4)
    set_table_borders(tbl_comp, color="CBD5E1", sz="4")
    
    comp_headers = ["Evaluation Criterion", "Traditional Ad-Hoc / Spreadsheet System", "Proposed PPMS Solution"]
    for i, h in enumerate(comp_headers):
        tbl_comp.cell(0, i).paragraphs[0].text = h
    format_row(tbl_comp.rows[0], "F1F5F9", is_header=True)
    
    comp_data = [
        ("Data Organization", "Disjointed text files or unconstrained spreadsheet cells.", "Strict Layered Architecture with domain encapsulation."),
        ("Product Identity", "Manual, error-prone IDs frequently duplicated or recycled.", "High-watermark atomic ID generator (PRD-xxxx); collision-free."),
        ("Pricing Lifecycle", "Direct in-place cell overwrites; no preview or approval.", "Dedicated workflow: preview diff & %, reason prompt, confirmation."),
        ("Audit Trail", "Non-existent; past prices and revision dates are lost.", "Permanent immutable PriceHistory ledger (HST-xxxx) with timestamps."),
        ("Financial Math", "Primitive float/double arithmetic subject to binary roundoff.", "java.math.BigDecimal with explicit Half-Up rounding policy."),
        ("Search & Filter", "Basic text search; rigid or non-deterministic filtering.", "Multi-attribute substring search; boundary-validated filtering."),
        ("Data Persistence", "Direct file writes vulnerable to mid-stream crash truncation.", "Atomic file replacement (.tmp -> .bak -> atomic rename)."),
        ("CSV Parsing", "Naive comma splitting; breaks on quotes and commas.", "RFC-4180 deterministic finite-state machine parser."),
        ("Executive Intelligence", "Requires manual formula authoring; no anomaly alerts.", "Live real-time dashboard: margins, medians, below-cost alerts.")
    ]
    for r_idx, (crit, trad, prop) in enumerate(comp_data):
        row = tbl_comp.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = crit
        row.cells[1].paragraphs[0].text = trad
        row.cells[2].paragraphs[0].text = prop
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # SECTION 7: FEASIBILITY STUDY
    # =========================================================================
    add_heading_1(doc, "7. System Feasibility Analysis")
    add_p(doc, "Prior to committing engineering resources, a comprehensive multi-dimensional feasibility study was conducted:")
    add_bullet(doc, "The project is exceptionally technically feasible. By restricting dependencies exclusively to Core Java SE 17, the application operates on any platform hosting an OpenJDK JRE. The standard libraries provide robust primitives for multithreading, file I/O, decimal arithmetic, and regular expressions.", bold_prefix="Technical Feasibility: ")
    add_bullet(doc, "A console-based interface is ideally suited for fast, keyboard-centric business workflows. Store managers and warehouse staff can navigate menus using single-digit choices and standard keybindings without requiring mouse navigation.", bold_prefix="Operational Feasibility: ")
    add_bullet(doc, "Economic feasibility is optimal. Development requires zero commercial software licenses, zero paid database servers, and zero cloud hosting subscriptions. The system runs entirely on existing standard workstation hardware.", bold_prefix="Economic Feasibility: ")
    add_bullet(doc, "The project scope was strategically structured into prioritized milestones (Core Catalog -> Pricing Governance -> Persistence -> Analytics -> Import/Export). This modular phasing guaranteed delivery of a working core before advanced features were introduced.", bold_prefix="Schedule Feasibility: ")
    add_bullet(doc, "The strict adherence to interface-driven design (ProductRepository, PriceHistoryRepository) ensures that underlying storage mechanisms can be maintained or upgraded with zero changes to business services.", bold_prefix="Maintainability Feasibility: ")

    # =========================================================================
    # SECTION 8: HARDWARE & SOFTWARE REQUIREMENTS
    # =========================================================================
    add_heading_1(doc, "8. Software & Hardware Engineering Requirements")
    add_p(doc, "The target execution profile for PPMS is deliberately lightweight, guaranteeing reliable deployment across legacy and modern workstations alike.")
    
    add_heading_2(doc, "8.1 Hardware Environment (Minimum & Recommended)")
    add_bullet(doc, "Processor: x86-64 / AMD64 dual-core processor operating at 1.5 GHz or higher (ARM64 supported via compatible OpenJDK).")
    add_bullet(doc, "Random Access Memory (RAM): Minimum 512 MB physical RAM; 1 GB or higher recommended.")
    add_bullet(doc, "Persistent Storage: 50 MB available hard disk space for Java binaries, CSV data repositories, and snapshot backups.")
    add_bullet(doc, "Display: Standard console terminal supporting minimum 80 columns by 24 lines (120 columns recommended for wide tabular layouts).")
    
    add_heading_2(doc, "8.2 Software Environment")
    add_bullet(doc, "Operating System: Microsoft Windows 10 / 11, Linux (Ubuntu, Debian, RHEL), or macOS 12+.")
    add_bullet(doc, "Java Runtime / Development Kit: OpenJDK 17 LTS (verified on Eclipse Temurin 17.0.20.1+1).")
    add_bullet(doc, "Terminal Shell: Windows PowerShell, Command Prompt (CMD), or Unix Bash.")
    add_bullet(doc, "Character Encoding: UTF-8 standard console encoding.")

    # =========================================================================
    # SECTION 9: SOFTWARE ENGINEERING PROCESS (ALL FIVE ACTIVITIES)
    # =========================================================================
    add_heading_1(doc, "9. Software Engineering Process: All Five Activities")
    add_p(doc, "In accordance with classical software engineering methodology, the development of PPMS was governed by the five fundamental generic process activities: Communication, Planning, Modelling, Construction, and Deployment.")

    add_heading_2(doc, "9.1 Activity 1: Communication")
    add_p(doc, "The Communication activity focused on customer and user requirement elicitation, domain exploration, ambiguity identification, and formal scope definition:")
    add_bullet(doc, "Identified primary actors as Retail Store Managers, Inventory Administrators, and Academic Evaluators. Defined high-level operational requirements for product maintenance and pricing governance.", bold_prefix="Stakeholder Identification & Elicitation: ")
    add_bullet(doc, "Resolved crucial domain ambiguities early: (1) Decided that products must be soft-archived rather than deleted to protect historical pricing integrity; (2) Established that zero-priced items are permitted only for promotional gifts/samples, whereas negative prices are strictly forbidden; (3) Handled zero old price in percentage change calculations by returning an empty indicator rather than causing an arithmetic crash.", bold_prefix="Ambiguity Resolution & Business Rules: ")
    add_bullet(doc, "Produced a formal Software Requirements Specification (SRS), problem statement, functional boundary matrix, and initial acceptance criteria.", bold_prefix="Communication Deliverables: ")

    add_heading_2(doc, "9.2 Activity 2: Planning")
    add_p(doc, "The Planning activity translated user requirements into an actionable technical roadmap, assessing risks, task dependencies, and resource allocations:")
    add_bullet(doc, "Divided implementation into five sequential milestones: Milestone 1 (Domain & Validation), Milestone 2 (Repositories & Persistence), Milestone 3 (Pricing Governance), Milestone 4 (Analytics & Reporting), Milestone 5 (Advanced Import/Export & Release).", bold_prefix="Work Breakdown & Milestones: ")
    add_bullet(doc, "Analyzed technical risks and formulated mitigation strategies: (1) CSV write corruption mitigated via atomic staging (.tmp) and backup (.bak); (2) Decimal roundoff errors eliminated by mandating BigDecimal; (3) Inconsistent terminal formatting mitigated by dynamic column width auto-calculation.", bold_prefix="Risk Management & Mitigation: ")
    add_bullet(doc, "Milestone execution schedule, technical dependency graph, automated test verification strategy, and academic team responsibility allocation matrix.", bold_prefix="Planning Deliverables: ")

    add_heading_2(doc, "9.3 Activity 3: Modelling")
    add_p(doc, "The Modelling activity transformed requirements into rigorous architectural, behavioral, structural, and data models prior to code construction:")
    add_bullet(doc, "Designed a 5-tier decoupled layered architecture separating Presentation, Service, Domain, Repository, and Persistence layers.", bold_prefix="Architectural Design: ")
    add_bullet(doc, "Engineered structural models for Product, PriceHistory, ProductStatus, and PriceProposal, defining strict encapsulation and immutability.", bold_prefix="Class & Domain Modelling: ")
    add_bullet(doc, "Drafted operational flowcharts for product creation, pricing proposal previews, confirmation workflows, search execution, and atomic CSV serialization.", bold_prefix="Behavioral & Process Flow Modelling: ")
    add_bullet(doc, "Designed normalized CSV schemas for products.csv, price_history.csv, and metadata.csv with strict data type specifications.", bold_prefix="Data & Schema Modelling: ")

    add_heading_2(doc, "9.4 Activity 4: Construction")
    add_p(doc, "The Construction activity focused on translating design models into robust, verified Core Java source code:")
    add_bullet(doc, "Implemented clean, encapsulated Java classes across com.ppms.model, com.ppms.service, com.ppms.repository, com.ppms.persistence, and com.ppms.ui.", bold_prefix="Component Coding: ")
    add_bullet(doc, "Engineered an RFC-4180 compliant finite-state machine CSV parser and atomic file replacement writer in CsvFileHandler.java.", bold_prefix="Persistence Engineering: ")
    add_bullet(doc, "Constructed a zero-dependency automated testing framework in Assert.java and TestRunner.java, validating 25 individual test cases.", bold_prefix="Automated Testing & QA: ")
    add_bullet(doc, "Cleanly compiling Java bytecode in bin/, modular source files, and 100% verified test execution reports.", bold_prefix="Construction Deliverables: ")

    add_heading_2(doc, "9.5 Activity 5: Deployment")
    add_p(doc, "The Deployment activity prepared the compiled software for immediate academic evaluation, operational execution, and demonstration:")
    add_bullet(doc, "Authored cross-platform compilation and execution scripts (compile.bat, run.bat, test.bat, build.ps1, run.ps1, compile.sh, run.sh).", bold_prefix="Build Automation: ")
    add_bullet(doc, "Engineered automatic sample data bootstrapping in Application.java, seeding 10 realistic products and historical price records upon first launch.", bold_prefix="Demonstration Bootstrapping: ")
    add_bullet(doc, "Validated clean execution from scratch on fresh workspace directories without requiring manual configuration.", bold_prefix="Release Verification: ")

    add_heading_2(doc, "9.6 Iterative Feedback Dynamics Among Activities")
    add_p(doc, "Although presented logically in sequential order, the engineering process incorporated critical iterative feedback loops. For example, during Construction unit testing of BackupService, it was discovered that two snapshots created within the same second collided under yyyyMMdd_HHmmss folder naming. This prompted a feedback loop to Modelling to refine the timestamp schema to yyyyMMdd_HHmmss_SSS with collision-safe loop guards, illustrating that real-world software engineering is dynamic and responsive to empirical findings.")

    # =========================================================================
    # SECTION 10: SYSTEM ARCHITECTURE & LOGICAL DESIGN
    # =========================================================================
    add_heading_1(doc, "10. System Architecture & Logical Design")
    add_p(doc, "PPMS is structured upon a strict, five-layer architectural model that guarantees complete separation of concerns and maintainability:")
    
    add_heading_2(doc, "10.1 Presentation Layer (com.ppms.ui, com.ppms.app)")
    add_bullet(doc, "Main / Application: System bootstrap, dependency injection wiring, directory verification, and master event loop initiation.")
    add_bullet(doc, "MenuController: Routes user menu selections across Catalog, Pricing, History, Analytics, Data Management, and Settings modules.")
    add_bullet(doc, "ConsoleUI: Encapsulates terminal I/O, keyboard parsing, input range checking, toggleable ANSI escape styling, and paginated navigation.")
    add_bullet(doc, "TableFormatter: Renders dynamic, perfectly aligned ASCII tables (+---, |) and Unicode box borders with auto-calculated column widths.")
    
    add_heading_2(doc, "10.2 Service Layer (com.ppms.service)")
    add_bullet(doc, "ProductService: Coordinates product registration, detail updates, soft-archival, restoration, multi-attribute searching, and deterministic sorting.")
    add_bullet(doc, "PricingService: Centralizes pricing governance; computes previews, diffs, and % changes; validates business rules; updates products; and logs audit records.")
    add_bullet(doc, "AnalyticsService: Computes real-time catalog metrics (median price, cost/retail valuations, gross margin %, markup %, below-cost alerts).")
    add_bullet(doc, "ImportExportService: Manages RFC-4180 CSV export and batch import with pre-validation and row-level rejection reporting.")
    add_bullet(doc, "BackupService: Creates timestamped snapshot archives and performs pre-validated restorations with safety rollback snapshots.")
    add_bullet(doc, "ReportGenerator: Formats executive valuation and category breakdown operational reports.")
    
    add_heading_2(doc, "10.3 Domain & Validation Layer (com.ppms.model, com.ppms.validation)")
    add_bullet(doc, "Product: Encapsulates product state, defensive copying, financial normalizations, and unit margin computations.")
    add_bullet(doc, "PriceHistory: Represents an immutable historical audit record capturing price changes, timestamps, and justifications.")
    add_bullet(doc, "PriceProposal: Value object encapsulating a proposed selling price modification prior to user confirmation.")
    add_bullet(doc, "ProductValidator: Centralized domain validation enforcing string lengths, ID patterns, non-negative quantities, and price rules.")
    
    add_heading_2(doc, "10.4 Repository Layer (com.ppms.repository)")
    add_bullet(doc, "ProductRepository: Interface defining catalog querying and persistence contracts, completely decoupling services from storage formats.")
    add_bullet(doc, "PriceHistoryRepository: Interface defining audit trail recording and querying contracts.")
    
    add_heading_2(doc, "10.5 Persistence Layer (com.ppms.persistence)")
    add_bullet(doc, "CsvProductRepository: In-memory cached implementation of ProductRepository with synchronized disk flushing.")
    add_bullet(doc, "CsvPriceHistoryRepository: In-memory cached implementation of PriceHistoryRepository.")
    add_bullet(doc, "CsvFileHandler: Coordinates atomic safe-writing (.tmp staging -> .bak backup -> atomic rename) and safe reading.")
    add_bullet(doc, "CsvParser: RFC-4180 compliant finite-state machine parsing commas, escaped quotes, and multiline cells.")
    add_bullet(doc, "CsvWriter: Formats and quotes cells containing special characters according to RFC-4180.")

    # =========================================================================
    # SECTION 11: DETAILED SYSTEM OPERATIONAL FLOWS
    # =========================================================================
    add_heading_1(doc, "11. Detailed System Operational Workflows")
    add_p(doc, "The operational workflows within PPMS are structured to enforce data integrity at every interaction boundary:")
    
    add_heading_2(doc, "11.1 Master Application Lifecycle Flow")
    add_p(doc, "1. Startup & Initialization: Main.java invokes Application.run(), which verifies local directory structures (data/ and data/backups/).\n2. High-Watermark Counter Sync: IdGenerator inspects existing records in products.csv and price_history.csv to initialize sequence counters.\n3. Demonstration Bootstrapping: If repositories are empty, 10 realistic products and 3 historical price changes are automatically seeded.\n4. Real-Time Dashboard Rendering: AnalyticsService calculates live catalog metrics and displays the master dashboard.\n5. Menu Loop & Input Routing: MenuController prompts the user for action selection, validating choices against permitted boundaries.\n6. Safe Termination: On user exit selection (0), all repositories verify disk synchronization before clean JVM shutdown.")
    
    add_heading_2(doc, "11.2 Product Creation Workflow")
    add_p(doc, "1. User initiates Add Product in Menu 1.\n2. Optional custom ID prompt: user inputs custom ID or leaves blank for automated sequence (PRD-xxxx).\n3. Input collection: Name, Category, Description, Cost Price, Selling Price, and Stock Quantity.\n4. Validation: ProductValidator checks string lengths, regex format, and non-negative values. ProductRepository verifies ID uniqueness.\n5. Rejection Branch: If validation fails, errors are displayed and no entity is created.\n6. Persistence Branch: If valid, Product is instantiated and saved to CsvProductRepository, updating cache and disk atomically.")
    
    add_heading_2(doc, "11.3 Selling Price Revision Workflow (Pricing Governance)")
    add_p(doc, "1. User selects Adjust Selling Price in Menu 2 and specifies Product ID.\n2. Product lookup: system displays current product name, cost price, and active selling price.\n3. Proposed price & reason input: user enters proposed selling price and mandatory business justification.\n4. Proposal generation: PricingService validates proposed price >= 0.00, verifies proposed != current price, and constructs PriceProposal.\n5. Tabular Preview Display: UI presents a structured preview showing Current Price, Proposed Price, Signed Difference, Percentage Change, and Cost Warnings.\n6. Confirmation Gateway: User is explicitly prompted 'Confirm price change? (yes/no)'.\n7. Cancellation Branch: If 'no', operation aborts immediately; product remains completely unchanged; zero history is logged.\n8. Approval Branch: If 'yes', Product selling price is updated and saved; a permanent PriceHistory record (HST-xxxx) is created and persisted atomically.")

    # =========================================================================
    # SECTION 12: USE-CASE MODEL & BEHAVIORAL SPECIFICATIONS
    # =========================================================================
    add_heading_1(doc, "12. Use-Case Model & Behavioral Specifications")
    add_p(doc, "The primary actor interacting with PPMS is the Store Manager / Pricing Administrator. Table 12.1 details the formal use-case specifications.")
    
    tbl_uc = doc.add_table(rows=9, cols=4)
    tbl_uc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_uc.autofit = False
    tbl_uc.columns[0].width = Inches(1.0)
    tbl_uc.columns[1].width = Inches(1.5)
    tbl_uc.columns[2].width = Inches(2.0)
    tbl_uc.columns[3].width = Inches(2.0)
    set_table_borders(tbl_uc, color="CBD5E1", sz="4")
    
    uc_headers = ["Use Case ID", "Use Case Name", "Trigger & Input", "System Response & Outcome"]
    for i, h in enumerate(uc_headers):
        tbl_uc.cell(0, i).paragraphs[0].text = h
    format_row(tbl_uc.rows[0], "F1F5F9", is_header=True)
    
    uc_data = [
        ("UC-01", "Register Product", "Manager selects Add Product; inputs attributes.", "Validates constraints; generates PRD-xxxx; persists to CSV."),
        ("UC-02", "Search & Filter", "Manager enters keyword or boundary criteria.", "Filters catalog in-memory; renders paginated tabular results."),
        ("UC-03", "Adjust Price", "Manager inputs Product ID, new price, and reason.", "Renders proposal preview; on confirmation, updates price & logs audit."),
        ("UC-04", "Bulk Price Revaluation", "Manager selects category or all products; enters %.", "Renders multi-item preview table; on confirmation, updates all items."),
        ("UC-05", "Audit History Query", "Manager requests history by product, date, or reason.", "Retrieves PriceHistory records; renders chronological audit log."),
        ("UC-06", "Generate Reports", "Manager selects Executive Valuation or Category report.", "AnalyticsService computes live statistics; renders formatted report."),
        ("UC-07", "CSV Batch Import", "Manager provides external CSV file path.", "Validates file; reports rejected rows by line; imports valid rows."),
        ("UC-08", "Backup & Restore", "Manager initiates snapshot backup or restore.", "Creates timestamped archive; on restore, takes safety backup & reloads.")
    ]
    for r_idx, (uid, uname, trig, outc) in enumerate(uc_data):
        row = tbl_uc.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = uid
        row.cells[1].paragraphs[0].text = uname
        row.cells[2].paragraphs[0].text = trig
        row.cells[3].paragraphs[0].text = outc
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # SECTION 13: UML CLASS MODEL & DATA DICTIONARY
    # =========================================================================
    add_heading_1(doc, "13. Class Model & CSV Data Dictionary")
    add_p(doc, "The structural architecture of PPMS is anchored upon clean, cohesive class definitions and well-documented file schemas:")
    
    add_heading_2(doc, "13.1 Core Class Specifications")
    add_bullet(doc, "Product (com.ppms.model): Represents a catalog entity. Attributes: productId (String), name (String), category (String), description (String), costPrice (BigDecimal), sellingPrice (BigDecimal), quantity (int), status (ProductStatus), createdAt (LocalDateTime), updatedAt (LocalDateTime). Methods: getGrossMarginAmount(), getGrossMarginPercentage(), getMarkupPercentage(), isBelowCost(), isLowStock(threshold), getInventoryCostValue(), getInventorySellingValue().")
    add_bullet(doc, "PriceHistory (com.ppms.model): Immutable audit entity. Attributes: historyId (String), productId (String), oldPrice (BigDecimal), newPrice (BigDecimal), changedAt (LocalDateTime), reason (String). Methods: getPriceDifference(), getPercentageChange(), isPriceIncrease(), isPriceDecrease().")
    add_bullet(doc, "PriceProposal (com.ppms.model): Value object for proposal previews. Attributes: product (Product), currentPrice (BigDecimal), proposedPrice (BigDecimal), reason (String), validationResult (ValidationResult).")
    
    add_heading_2(doc, "13.2 CSV Schema Data Dictionary")
    add_p(doc, "Table 13.1 defines the schema for products.csv:")
    
    tbl_p_sch = doc.add_table(rows=11, cols=4)
    tbl_p_sch.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_p_sch.autofit = False
    tbl_p_sch.columns[0].width = Inches(1.3)
    tbl_p_sch.columns[1].width = Inches(1.2)
    tbl_p_sch.columns[2].width = Inches(1.5)
    tbl_p_sch.columns[3].width = Inches(2.5)
    set_table_borders(tbl_p_sch, color="CBD5E1", sz="4")
    
    p_headers = ["Field Name", "Data Type", "Constraints / Format", "Description"]
    for i, h in enumerate(p_headers):
        tbl_p_sch.cell(0, i).paragraphs[0].text = h
    format_row(tbl_p_sch.rows[0], "F1F5F9", is_header=True)
    
    p_fields = [
        ("productId", "String", "Unique, Non-Null, PRD-xxxx", "Permanent alphanumeric primary identifier."),
        ("name", "String", "2-120 chars, Non-Blank", "Commercial name of the product."),
        ("category", "String", "2-60 chars, Non-Blank", "Merchandise department/category classification."),
        ("description", "String", "0-500 chars, Optional", "Commercial description; quoted if contains commas."),
        ("costPrice", "BigDecimal", ">= 0.00, Scale 2", "Wholesale unit procurement cost."),
        ("sellingPrice", "BigDecimal", ">= 0.00, Scale 2", "Active retail selling price."),
        ("quantity", "Integer", ">= 0, Non-Negative", "Physical inventory units currently in stock."),
        ("status", "Enum", "ACTIVE, ARCHIVED", "Product lifecycle state for soft-archival."),
        ("createdAt", "DateTime", "ISO-8601 (yyyy-MM-dd'T'HH:mm:ss)", "Timestamp of initial registration."),
        ("updatedAt", "DateTime", "ISO-8601 (yyyy-MM-dd'T'HH:mm:ss)", "Timestamp of most recent modification.")
    ]
    for r_idx, (fn, dt, con, desc) in enumerate(p_fields):
        row = tbl_p_sch.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = fn
        row.cells[1].paragraphs[0].text = dt
        row.cells[2].paragraphs[0].text = con
        row.cells[3].paragraphs[0].text = desc
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # SECTION 14: FORMAL FUNCTIONAL REQUIREMENTS
    # =========================================================================
    add_heading_1(doc, "14. Formal Functional Requirements")
    add_p(doc, "Table 14.1 documents the complete functional requirements matrix implemented within PPMS.")
    
    tbl_fr = doc.add_table(rows=17, cols=4)
    tbl_fr.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_fr.autofit = False
    tbl_fr.columns[0].width = Inches(0.9)
    tbl_fr.columns[1].width = Inches(1.3)
    tbl_fr.columns[2].width = Inches(3.3)
    tbl_fr.columns[3].width = Inches(1.0)
    set_table_borders(tbl_fr, color="CBD5E1", sz="4")
    
    fr_headers = ["Req ID", "Module", "Functional Requirement Description", "Priority"]
    for i, h in enumerate(fr_headers):
        tbl_fr.cell(0, i).paragraphs[0].text = h
    format_row(tbl_fr.rows[0], "F1F5F9", is_header=True)
    
    fr_data = [
        ("FR-01", "Catalog", "Register new products with automated or custom unique IDs.", "Mandatory"),
        ("FR-02", "Catalog", "View active and archived products in paginated tabular format.", "Mandatory"),
        ("FR-03", "Catalog", "Search products across ID, name, category, and description.", "Mandatory"),
        ("FR-04", "Catalog", "Update non-price details (name, category, description).", "Mandatory"),
        ("FR-05", "Catalog", "Soft-archive and restore products without destroying history.", "Mandatory"),
        ("FR-06", "Pricing", "Update selling price with preview of diff, % change, and reason.", "Mandatory"),
        ("FR-07", "Audit", "Log immutable PriceHistory records on every approved price change.", "Mandatory"),
        ("FR-08", "Audit", "Filter audit history by product ID, date range, and reason keyword.", "High"),
        ("FR-09", "Audit", "Identify historical largest price increase and largest price decrease.", "High"),
        ("FR-10", "Catalog", "Multi-criteria filtering with price/quantity boundary validation.", "High"),
        ("FR-11", "Catalog", "Deterministic sorting by name, price, cost, quantity, or date.", "High"),
        ("FR-12", "Analytics", "Calculate live metrics: min, max, avg, median, cost/retail values.", "High"),
        ("FR-13", "Diagnostics", "Identify loss-leader products (selling < cost) and low-stock items.", "High"),
        ("FR-14", "Interoperability", "Export products and audit trail to RFC-4180 CSV files.", "Medium"),
        ("FR-15", "Interoperability", "Batch import products with row-by-row pre-validation & rejection log.", "Medium"),
        ("FR-16", "Data Recovery", "Create timestamped snapshot backups and restore with safety backup.", "Medium")
    ]
    for r_idx, (rid, mod, desc, prio) in enumerate(fr_data):
        row = tbl_fr.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = rid
        row.cells[1].paragraphs[0].text = mod
        row.cells[2].paragraphs[0].text = desc
        row.cells[3].paragraphs[0].text = prio
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # SECTION 15: NON-FUNCTIONAL REQUIREMENTS
    # =========================================================================
    add_heading_1(doc, "15. Non-Functional Requirements & Engineering Standards")
    add_bullet(doc, "Financial Arithmetic Rigor: Zero usage of primitive float or double for monetary attributes. All currency storage and arithmetic strictly utilize java.math.BigDecimal with explicit Half-Up rounding, guaranteeing 100% financial calculation accuracy.", bold_prefix="NFR-01 (Financial Rigor): ")
    add_bullet(doc, "Crash-Safe Persistence: System writes updates to a temporary staging file (.tmp), flushes to disk, creates a backup (.bak), and executes an atomic filesystem rename, ensuring zero partial-write corruption.", bold_prefix="NFR-02 (Data Integrity): ")
    add_bullet(doc, "High-Watermark Collision Protection: Unique sequential IDs (PRD-1001, HST-1001) are synchronized with existing CSV records on startup, guaranteeing that IDs are never recycled or duplicated across restarts.", bold_prefix="NFR-03 (Identifier Permanence): ")
    add_bullet(doc, "Universal Cross-Platform Portability: Written entirely in standard Core Java SE 17 without JNI or platform-specific commands, executing identically on Windows, Linux, and macOS.", bold_prefix="NFR-04 (Portability): ")
    add_bullet(doc, "Terminal UI Usability: Structured visual design featuring clean ASCII tables (+---, |), dynamic column width formatting, pagination for long records, and user confirmations for destructive operations.", bold_prefix="NFR-05 (Usability): ")
    add_bullet(doc, "Autonomous Offline Execution: Operates completely disconnected from the internet, requiring zero external database installations or cloud dependencies.", bold_prefix="NFR-06 (Self-Sufficiency): ")

    # =========================================================================
    # SECTION 16: VERIFICATION & TEST EXECUTION REPORT
    # =========================================================================
    add_heading_1(doc, "16. Verification, Quality Assurance & Test Execution Report")
    add_p(doc, "Quality assurance was executed using a dedicated, zero-dependency automated testing engine implemented in Assert.java and coordinated by TestRunner.java. Table 16.1 summarizes the verified test results across all 6 test suites.")
    
    tbl_ts = doc.add_table(rows=7, cols=5)
    tbl_ts.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ts.autofit = False
    tbl_ts.columns[0].width = Inches(1.8)
    tbl_ts.columns[1].width = Inches(2.2)
    tbl_ts.columns[2].width = Inches(0.8)
    tbl_ts.columns[3].width = Inches(0.8)
    tbl_ts.columns[4].width = Inches(0.9)
    set_table_borders(tbl_ts, color="CBD5E1", sz="4")
    
    ts_headers = ["Test Suite Class", "Target Focus & Verification Scope", "Tests", "Passed", "Pass Rate"]
    for i, h in enumerate(ts_headers):
        tbl_ts.cell(0, i).paragraphs[0].text = h
    format_row(tbl_ts.rows[0], "F1F5F9", is_header=True)
    
    ts_data = [
        ("ProductTest", "Entity validation, duplicate ID rejection, invalid prices, soft-archival, search & filter.", "8", "8", "100.0%"),
        ("PricingTest", "Price proposal preview, signed diff, % change formula, identical price rejection, zero old price.", "6", "6", "100.0%"),
        ("PersistenceTest", "RFC-4180 quote escaping, restart persistence, malformed row resilience, snapshot backup/restore.", "5", "5", "100.0%"),
        ("AnalyticsTest", "Median calculation, cost/retail valuations, gross margin %, markup %, empty dataset safety.", "3", "3", "100.0%"),
        ("ImportExportTest", "Batch CSV import, duplicate rejection, line-by-line error log, round-trip export verification.", "2", "2", "100.0%"),
        ("EndToEndTest", "Complete 11-step lifecycle: Add -> Search -> Update -> Audit -> Analytics -> Export -> Restart Survival.", "1", "1", "100.0%")
    ]
    for r_idx, (sc, foc, tst, pas, prt) in enumerate(ts_data):
        row = tbl_ts.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = sc
        row.cells[1].paragraphs[0].text = foc
        row.cells[2].paragraphs[0].text = tst
        row.cells[3].paragraphs[0].text = pas
        row.cells[4].paragraphs[0].text = prt
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    add_p(doc, "Overall Testing Verdict: 6 of 6 Test Suites Passed (25 of 25 Test Cases Verified; 100% Pass Rate).", bold_prefix="Automated QA Conclusion: ")

    # =========================================================================
    # SECTION 17: CONSOLE USER INTERFACE DESIGN
    # =========================================================================
    add_heading_1(doc, "17. Console User Interface Design & Terminal UX")
    add_p(doc, "The presentation layer was deliberately crafted to transcend the look of beginner-level print statements, delivering a structured, professional console management interface.")
    add_heading_2(doc, "17.1 Visual Direction & Typography")
    add_bullet(doc, "Border Typography: Standardized clean ASCII (+---, |) borders by default, ensuring 100% universal compatibility across legacy Windows Command Prompt, PowerShell, and Unix shells without '?' character corruption.")
    add_bullet(doc, "Restrained Color Palette: Cyan and bold blue for application headers and section markers; green [SUCCESS] for confirmed operations; yellow [WARNING] for diagnostic alerts; red [ERROR] for rejected input.")
    add_bullet(doc, "Consistent Currency Formatting: Normalized financial representation (Rs. 1,499.00 or ₹1,499.00) with comma grouping separators and 2 decimal places.")
    add_bullet(doc, "Pagination Engine: Tables exceeding page limits are presented in paginated views (Page X of Y) with [N]ext, [P]revious, and [G]o-to controls.")
    
    add_heading_2(doc, "17.2 Representative Terminal Output")
    add_p(doc, "Figure 17.1 illustrates the live rendered application header and real-time dashboard:")
    add_callout(doc, "Live Dashboard Terminal Render",
"+---------------------------+-------+---------------------------+---------------+\n"
"|          METRIC           | VALUE |          METRIC           |     VALUE     |\n"
"+---------------------------+-------+---------------------------+---------------+\n"
"| Total Registered Products | 10    | Lowest Active Price       | Rs. 99.00     |\n"
"| Active Products           | 10    | Highest Active Price      | Rs. 54,990.00 |\n"
"| Archived Products         | 0     | Average Active Price      | Rs. 11,946.80 |\n"
"| Total Stock Units         | 383   | Median Active Price       | Rs. 5,999.00  |\n"
"| Priced Below Cost [ALERT] | 2     | Total Inventory Valuation | Rs. 2,050,194 |\n"
"| Low Stock Items (<= 10)   | 2     | Price Changes Logged      | 3             |\n"
"+---------------------------+-------+---------------------------+---------------+\n\n"
"Recent Price Adjustments:\n"
"  * [2026-10-07 12:29:39] Product PRD-1001: Rs. 8,495.00 -> Rs. 8,995.00 (+5.89%) | Supplier cost increase\n"
"  * [2026-10-15 12:29:39] Product PRD-1002: Rs. 58,990.00 -> Rs. 54,990.00 (-6.78%) | Festive discount\n"
"  * [2026-10-19 12:29:39] Product PRD-1007: Rs. 32,990.00 -> Rs. 29,990.00 (-9.09%) | Competitor matching",
        bg_color="F1F5F9", border_color="1B365D"
    )

    # =========================================================================
    # SECTION 18: PROJECT IMPLEMENTATION PLAN & MILESTONES
    # =========================================================================
    add_heading_1(doc, "18. Project Implementation Plan & Milestones")
    add_p(doc, "Table 18.1 details the engineering milestone roadmap followed during the development of PPMS.")
    
    tbl_ms = doc.add_table(rows=6, cols=4)
    tbl_ms.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ms.autofit = False
    tbl_ms.columns[0].width = Inches(1.1)
    tbl_ms.columns[1].width = Inches(1.8)
    tbl_ms.columns[2].width = Inches(2.3)
    tbl_ms.columns[3].width = Inches(1.3)
    set_table_borders(tbl_ms, color="CBD5E1", sz="4")
    
    ms_headers = ["Milestone", "Target Module", "Key Engineering Deliverables", "Status"]
    for i, h in enumerate(ms_headers):
        tbl_ms.cell(0, i).paragraphs[0].text = h
    format_row(tbl_ms.rows[0], "F1F5F9", is_header=True)
    
    ms_data = [
        ("Milestone 1", "Domain & Validation", "Product.java, ProductStatus.java, ProductValidator.java, ValidationResult.java.", "Completed"),
        ("Milestone 2", "Repositories & Persistence", "CsvProductRepository.java, CsvFileHandler.java, CsvParser.java, atomic rename.", "Completed"),
        ("Milestone 3", "Pricing & Audit Trail", "PricingService.java, PriceHistory.java, PriceProposal.java, preview workflow.", "Completed"),
        ("Milestone 4", "UI & Analytics Engine", "ConsoleUI.java, MenuController.java, TableFormatter.java, AnalyticsService.java.", "Completed"),
        ("Milestone 5", "Advanced & Release", "ImportExportService.java, BackupService.java, ReportGenerator.java, TestRunner.java.", "Completed")
    ]
    for r_idx, (ms, mod, deliv, st) in enumerate(ms_data):
        row = tbl_ms.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = ms
        row.cells[1].paragraphs[0].text = mod
        row.cells[2].paragraphs[0].text = deliv
        row.cells[3].paragraphs[0].text = st
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # SECTION 19: TEAM OWNERSHIP & RESPONSIBILITY MATRIX
    # =========================================================================
    add_heading_1(doc, "19. Academic Team Ownership & Responsibility Matrix")
    add_p(doc, "In accordance with university engineering evaluation standards, Table 19.1 presents the formal academic responsibility allocation across the six registered team members. Candidate registration numbers are preserved exactly as provided in university administrative records.")
    
    tbl_tm = doc.add_table(rows=7, cols=4)
    tbl_tm.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_tm.autofit = False
    tbl_tm.columns[0].width = Inches(1.1)
    tbl_tm.columns[1].width = Inches(1.4)
    tbl_tm.columns[2].width = Inches(1.8)
    tbl_tm.columns[3].width = Inches(2.2)
    set_table_borders(tbl_tm, color="CBD5E1", sz="4")
    
    tm_headers = ["Member", "Registration No.", "Assigned Role", "Scope & Deliverables"]
    for i, h in enumerate(tm_headers):
        tbl_tm.cell(0, i).paragraphs[0].text = h
    format_row(tbl_tm.rows[0], "F1F5F9", is_header=True)
    
    tm_data = [
        ("Team Member 1", "2025DCSBT10650", "Project Lead & Architect", "System architecture, layered design, dependency injection, and release."),
        ("Team Member 2", "2025DCSBT10651", "Product Model & Validation", "Product entity design, ProductValidator, BigDecimal financial normalization."),
        ("Team Member 3", "2025DCSBT10652", "Product Operations & Search", "ProductService, multi-attribute searching, deterministic sorting, filtering."),
        ("Team Member 4", "2025DCSBT10653", "Pricing & Historical Audit", "PricingService, PriceHistory ledger, proposal previews, percentage math."),
        ("Team Member 5", "2025DCSBT10654", "Persistence & QA Engineer", "RFC-4180 CSV parser/writer, atomic safe-write mechanics, TestRunner suite."),
        ("Team Member 6", "2025DCSBT10655", "Documentation & Demonstration", "ConsoleUI layout design, TableFormatter, user manual, and viva preparation.")
    ]
    for r_idx, (mem, reg, rol, scp) in enumerate(tm_data):
        row = tbl_tm.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = mem
        row.cells[1].paragraphs[0].text = reg
        row.cells[2].paragraphs[0].text = rol
        row.cells[3].paragraphs[0].text = scp
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # SECTION 20: SYSTEM RESULTS, CONSTRAINTS & ROADMAP
    # =========================================================================
    add_heading_1(doc, "20. System Results, Limitations & Strategic Roadmap")
    add_p(doc, "Results Summary: The Product Price Management System successfully delivers 100% of functional requirements FR-01 through FR-16. It operates autonomously in pure Core Java SE 17, persists data reliably across application restarts, maintains an auditable pricing trail, and executes clean verification workflows.")
    
    add_heading_2(doc, "20.1 Technical Limitations")
    add_bullet(doc, "Concurrent Multi-User Contention: Because persistence utilizes local flat CSV files, concurrent multi-user write operations across independent OS processes without an external lock manager could lead to file locking conflicts.", bold_prefix="Local File Concurrency: ")
    add_bullet(doc, "Relational Indexing: Substring search and filtering are executed in-memory via Java Streams. While highly performant for catalogs up to 50,000 items, extremely large datasets would benefit from B-Tree database indexing.", bold_prefix="In-Memory Querying: ")
    
    add_heading_2(doc, "20.2 Strategic Future Roadmap")
    add_bullet(doc, "Role-Based Access Control (RBAC): Integrate cryptographic password hashing (PBKDF2/BCrypt) in Core Java to distinguish Store Clerk vs Commercial Manager privileges.")
    add_bullet(doc, "Algorithmic Dynamic Markdown Pricing: Implement automated price markdown schedules triggered when stock levels remain stagnant past configurable aging thresholds.")
    add_bullet(doc, "JavaFX Desktop GUI Companion: Develop an optional graphical dashboard utilizing the existing service and repository interfaces without altering underlying business rules.")

    # =========================================================================
    # SECTION 21: CONCLUSION
    # =========================================================================
    add_heading_1(doc, "21. Conclusion")
    add_p(doc, "The Product Price Management System (PPMS) demonstrates the successful application of rigorous software engineering principles to commercial inventory pricing governance. By combining a strictly decoupled five-layer Object-Oriented architecture, crash-resilient RFC-4180 CSV persistence, precision BigDecimal financial arithmetic, and a highly usable interactive terminal interface, PPMS bridges the critical divide between simple catalog recording and enterprise pricing intelligence.")
    add_p(doc, "The project stands as proof that robust, auditable, and production-grade software does not inherently mandate heavyweight web frameworks or external database infrastructure. Through disciplined software engineering—spanning communication, planning, modelling, construction, and deployment—PPMS provides an educational and practical exemplar of Core Java software excellence.")

    # =========================================================================
    # SECTION 22: REFERENCES
    # =========================================================================
    add_heading_1(doc, "22. References & Academic Bibliography")
    add_bullet(doc, "Oracle Corporation. (2026). Java Platform, Standard Edition Documentation (JDK 17). Oracle Technical Network. https://docs.oracle.com/en/java/javase/17/")
    add_bullet(doc, "Bloch, J. (2018). Effective Java (3rd Edition). Addison-Wesley Professional.")
    add_bullet(doc, "Pressman, R. S., & Maxim, B. R. (2020). Software Engineering: A Practitioner's Approach (9th Edition). McGraw-Hill Education.")
    add_bullet(doc, "Shafranovich, Y. (2005). Common Format and MIME Type for Comma-Separated Values (CSV) Files (RFC 4180). Internet Engineering Task Force (IETF). https://www.ietf.org/rfc/rfc4180.txt")
    add_bullet(doc, "Gamma, E., Helm, R., Johnson, R., & Vlissides, J. (1994). Design Patterns: Elements of Reusable Object-Oriented Software. Addison-Wesley.")

    # =========================================================================
    # SECTION 23: APPENDIX — VIVA VOCE DEFENSE & TECHNICAL Q&A
    # =========================================================================
    add_heading_1(doc, "23. Appendix: Viva Voce Defense & Technical Examination Guide")
    add_p(doc, "This section prepares the academic candidate for oral examination and technical defense before a faculty evaluation committee:")
    
    add_heading_2(doc, "Q1: Why was java.math.BigDecimal selected over double for financial arithmetic?")
    add_p(doc, "Answer: Primitive double and float types rely on binary floating-point representation (IEEE-754), which cannot accurately represent decimal fractions like 0.1 or 0.05. In commercial pricing, cumulative arithmetic introduces inexact results (e.g. 0.1 + 0.2 = 0.30000000000000004). BigDecimal stores arbitrary-precision signed decimal numbers and provides explicit rounding modes (RoundingMode.HALF_UP), guaranteeing 100% mathematical precision in commercial audits.")
    
    add_heading_2(doc, "Q2: Why was local CSV persistence selected rather than a relational database (e.g. MySQL or SQLite)?")
    add_p(doc, "Answer: The project specification strictly mandated Core Java without external database drivers, ORMs, or native server installations. CSV provides universal cross-platform portability, human inspection, zero-install deployment, and offline self-sufficiency. Reliability is achieved through atomic file staging (.tmp) and backup (.bak) renaming.")
    
    add_heading_2(doc, "Q3: How does the application prevent data loss during unexpected crashes while writing files?")
    add_p(doc, "Answer: Direct file overwriting is prohibited in PPMS. Instead, CsvFileHandler writes new records to a staging file (products.csv.tmp). The stream is flushed and synced to disk. If the live file exists, it is backed up to products.csv.bak. Then, Files.move() executes an atomic filesystem rename (ATOMIC_MOVE, falling back to REPLACE_EXISTING). In-memory caches are updated only after the disk operation completes successfully.")
    
    add_heading_2(doc, "Q4: Why is String.split(\",\") unacceptable for professional CSV parsing?")
    add_p(doc, "Answer: Real-world merchandise descriptions frequently contain embedded commas (e.g., 'Logitech MX Master 3S, Wireless'), quotation marks, or newlines. Naive string splitting corrupts column alignment. PPMS implements an RFC-4180 compliant finite-state machine that tracks character-by-character quote states, escaped quote pairs (\"\"), and multiline records seamlessly.")
    
    add_heading_2(doc, "Q5: How is price change percentage computed when the previous price was zero?")
    add_p(doc, "Answer: The standard mathematical formula is ((NewPrice - OldPrice) / OldPrice) * 100. When OldPrice = 0.00, dividing by zero is mathematically undefined. PPMS prevents ArithmeticException by returning Optional.empty(), cleanly rendered in the console UI as 'N/A'.")
    
    add_heading_2(doc, "Q6: How does the Repository Pattern benefit the architectural maintainability of PPMS?")
    add_p(doc, "Answer: The Repository Pattern decouples domain and service logic from data persistence mechanisms. ProductService and PricingService depend exclusively on the ProductRepository and PriceHistoryRepository interfaces, never on concrete CSV classes. If storage is migrated to an encrypted file or database in the future, services remain 100% untouched.")

    # Save document
    output_path = os.path.abspath("Product_Price_Management_System_Project_Report.docx")
    doc.save(output_path)
    print(f"[SUCCESS] Report successfully generated at: {output_path}")

if __name__ == "__main__":
    build_docx()
