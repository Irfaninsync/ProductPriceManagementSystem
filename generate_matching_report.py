import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:right w:val="none"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>')
        tblPr[0].append(borders)

def format_row(row, bg_hex, is_header=False):
    for cell in row.cells:
        if bg_hex:
            set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.5 if not is_header else 10)
                if is_header:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(27, 54, 93) # Deep Navy
                else:
                    run.font.color.rgb = RGBColor(45, 55, 72) # Charcoal

def add_header(doc, text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    for r in h.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(27, 54, 93)
    return h

def add_sub_header(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    for r in h.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(43, 108, 176)
    return h

def add_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(27, 54, 93)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(45, 55, 72)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(45, 55, 72)
    return p

def add_box_diagram(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="12" w:space="0" w:color="2B6CB0"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(2)
    cp.paragraph_format.line_spacing = 1.05
    r = cp.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(30, 41, 59)

def generate_report():
    doc = Document()
    
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # =========================================================================
    # TITLE & ROLL NUMBERS AT THE TOP
    # =========================================================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Product Price Management System (PPMS)")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Core Java Enterprise Console Application | Local CSV Persistence")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(74, 85, 104)

    # Roll Numbers Table
    tbl_rolls = doc.add_table(rows=7, cols=3)
    tbl_rolls.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_rolls.autofit = False
    tbl_rolls.columns[0].width = Inches(1.8)
    tbl_rolls.columns[1].width = Inches(2.3)
    tbl_rolls.columns[2].width = Inches(2.7)
    set_table_borders(tbl_rolls, color="CBD5E1", sz="4")

    roll_headers = ["Candidate Roll Number", "Team Member Identifier", "Project Responsibility Area"]
    for i, h in enumerate(roll_headers):
        tbl_rolls.cell(0, i).paragraphs[0].text = h
    format_row(tbl_rolls.rows[0], "F1F5F9", is_header=True)

    rolls_data = [
        ("2025DCSBT10650", "Team Member 1", "Project Lead & System Architect"),
        ("2025DCSBT10651", "Team Member 2", "Product Model & Data Validation"),
        ("2025DCSBT10652", "Team Member 3", "Product Operations & Search Engine"),
        ("2025DCSBT10653", "Team Member 4", "Pricing Operations & Historical Audit"),
        ("2025DCSBT10654", "Team Member 5", "CSV Persistence & Automated Testing"),
        ("2025DCSBT10655", "Team Member 6", "Documentation & Console UI/UX Design")
    ]
    for r_idx, (roll, mem, resp) in enumerate(rolls_data):
        row = tbl_rolls.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = roll
        row.cells[1].paragraphs[0].text = mem
        row.cells[2].paragraphs[0].text = resp
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # 1. PROJECT OVERVIEW
    # =========================================================================
    add_header(doc, "1. Project Overview")
    add_p(doc, "A Product Price Management System (PPMS) is an advanced desktop console software application that stores, governs, and analyzes commercial product information and pricing lifecycles using local file storage.")
    add_p(doc, "The system enables small-to-medium retail businesses, store managers, pricing analysts, and inventory administrators to manage product catalogs, enforce formal price-change approvals, preserve an immutable audit trail of every historical price modification, analyze gross margins, identify pricing anomalies, and maintain persistent local records completely offline without requiring an internet connection or external database server.")

    # =========================================================================
    # 2. PROBLEM STATEMENT
    # =========================================================================
    add_header(doc, "2. Problem Statement")
    add_p(doc, "Traditional retail and catalog management commonly involves manual paper ledgers, standalone spreadsheets, or basic single-table software that store only current prices. This leads to:")
    add_bullet(doc, "Permanent loss of historical pricing context whenever prices are directly overwritten.")
    add_bullet(doc, "Absence of operational audit trails explaining why prices changed or who authorized modifications.")
    add_bullet(doc, "Inaccurate financial calculations caused by floating-point binary roundoff errors in standard programs.")
    add_bullet(doc, "Data corruption and truncation risks during file writes or unexpected system power loss.")
    add_bullet(doc, "Difficulty in searching, filtering, and sorting products across large inventories.")
    add_bullet(doc, "Lack of proactive alerts for loss-leader products sold below cost price or depleted stock levels.")
    add_bullet(doc, "Dependency on heavy, complex database infrastructure that is costly and difficult to maintain offline.")

    # =========================================================================
    # 3. PROPOSED SOLUTION
    # =========================================================================
    add_header(doc, "3. Proposed Solution")
    add_p(doc, "The proposed system provides an autonomous, decoupled Core Java platform for complete product and pricing management.")
    add_p(doc, "Main modules:", bold_prefix=None)
    
    add_sub_header(doc, "Product Catalog Management")
    add_bullet(doc, "Register new products with automated high-watermark unique IDs (PRD-xxxx)")
    add_bullet(doc, "Maintain product profiles (Name, Category, Description, Cost, Selling Price, Quantity)")
    add_bullet(doc, "Multi-attribute search across ID, name, category, and description")
    add_bullet(doc, "Deterministic sorting and multi-boundary filtering")
    add_bullet(doc, "Soft product archival and restoration preserving complete historical data")
    
    add_sub_header(doc, "Pricing Management")
    add_bullet(doc, "Inspect current selling price and unit margins")
    add_bullet(doc, "Calculate proposed price difference and percentage price change")
    add_bullet(doc, "Require mandatory business justification for every price update")
    add_bullet(doc, "Interactive pre-commit preview with explicit confirmation gateway")
    add_bullet(doc, "Single-item adjustments and category-wide bulk percentage adjustments")
    
    add_sub_header(doc, "Price History & Audit Trail")
    add_bullet(doc, "Log immutable audit records (historyId, productId, oldPrice, newPrice, timestamp, reason)")
    add_bullet(doc, "Retrieve chronological price history for any specific product")
    add_bullet(doc, "Filter audit records by date range and reason keyword")
    add_bullet(doc, "Identify historical extremes (largest price increase and largest price decrease)")
    
    add_sub_header(doc, "Pricing Analytics & Valuation")
    add_bullet(doc, "Calculate lowest, highest, average, and median active selling prices")
    add_bullet(doc, "Compute total stock units, inventory cost value, and inventory retail selling valuation")
    add_bullet(doc, "Compute gross margin amounts, gross margin percentage, and markup percentage")
    add_bullet(doc, "Flag loss-leader products priced below procurement cost (Selling < Cost)")
    add_bullet(doc, "Identify low-stock items falling at or below a configurable threshold")
    
    add_sub_header(doc, "Local CSV Persistence & Recovery")
    add_bullet(doc, "Store records locally in normalized UTF-8 CSV files (products.csv, price_history.csv, metadata.csv)")
    add_bullet(doc, "RFC-4180 compliant finite-state machine parser handling embedded commas, quotes, and newlines")
    add_bullet(doc, "Crash-safe atomic file writes (.tmp staging -> .bak backup -> atomic filesystem rename)")
    add_bullet(doc, "Create timestamped snapshot backups and restore with pre-validation and safety backups")
    
    add_sub_header(doc, "Admin Dashboard & Console UI")
    add_bullet(doc, "Real-time master dashboard displaying 12 operational statistics and recent price events")
    add_bullet(doc, "Keyboard-driven menus, paginated record navigation, and clean ASCII border tables")
    add_bullet(doc, "System settings configuration (Currency symbol, low-stock threshold, ANSI styling)")

    # =========================================================================
    # 4. SYSTEM ARCHITECTURE
    # =========================================================================
    add_header(doc, "4. System Architecture")
    add_p(doc, "PPMS is built upon a strict, decoupled Layered Object-Oriented Architecture:")
    
    add_box_diagram(doc,
"+-------------------------------------------------------------------------+\n"
"|                           Store Administrator                           |\n"
"+-------------------------------------------------------------------------+\n"
"                                     |\n"
"                                     v\n"
"+-------------------------------------------------------------------------+\n"
"|                    Presentation Layer (Console UI)                      |\n"
"|  Main Entry | MenuController | ConsoleUI (ANSI/Borders) | TableFormatter|\n"
"+-------------------------------------------------------------------------+\n"
"                                     |\n"
"                                     v\n"
"+-------------------------------------------------------------------------+\n"
"|                       Business Service Layer                            |\n"
"|  ProductService | PricingService | AnalyticsService | ImportExport      |\n"
"+-------------------------------------------------------------------------+\n"
"                    |                                 |\n"
"                    v                                 v\n"
"+---------------------------------------+  +------------------------------+\n"
"|       Domain & Validation Layer       |  |      Repository Layer        |\n"
"|  Product | PriceHistory | Validator   |  |  ProductRepo | PriceHistRepo |\n"
"+---------------------------------------+  +------------------------------+\n"
"                                                          |\n"
"                                                          v\n"
"+-------------------------------------------------------------------------+\n"
"|               Persistence Layer (CSV Storage Engine)                    |\n"
"|  CsvProductRepository | CsvPriceHistoryRepo | CsvFileHandler (Atomic)   |\n"
"|  RFC-4180 CsvParser (State Machine) | CsvWriter (Escaper)               |\n"
"+-------------------------------------------------------------------------+\n"
"                                     |\n"
"                                     v\n"
"+-------------------------------------------------------------------------+\n"
"|                     Local File System Storage                           |\n"
"|  data/products.csv | data/price_history.csv | data/metadata.csv | backups|\n"
"+-------------------------------------------------------------------------+"
    )

    # =========================================================================
    # 5. SOFTWARE ENGINEERING PROCESS FLOW (IN THE MIDDLE)
    # =========================================================================
    add_header(doc, "5. Software Engineering Process Flow")
    add_p(doc, "The development of PPMS was executed using the disciplined Five Generic Software Engineering Activities, structured with continuous quality feedback loops:")

    add_box_diagram(doc,
"[1. COMMUNICATION]  Requirement Elicitation, Stakeholder Goals, Constraint Definition\n"
"        |\n"
"        v\n"
"[2. PLANNING]       Milestone Scheduling, Risk Analysis, Architecture Dependency Graph\n"
"        |\n"
"        v\n"
"[3. MODELLING]      Domain Models, Layered Architecture, Flowcharts, CSV Schemas\n"
"        |\n"
"        v\n"
"[4. CONSTRUCTION]   Core Java SE 17 Coding, BigDecimal Arithmetic, Automated Testing\n"
"        |\n"
"        v\n"
"[5. DEPLOYMENT]     Build Scripts, Demonstration Bootstrapping, Clean Launch Verification"
    )

    tbl_se = doc.add_table(rows=6, cols=4)
    tbl_se.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_se.autofit = False
    tbl_se.columns[0].width = Inches(1.3)
    tbl_se.columns[1].width = Inches(1.8)
    tbl_se.columns[2].width = Inches(2.2)
    tbl_se.columns[3].width = Inches(1.5)
    set_table_borders(tbl_se, color="CBD5E1", sz="4")

    se_headers = ["SE Activity", "Core Objective", "PPMS Project Implementation", "Key Deliverables"]
    for i, h in enumerate(se_headers):
        tbl_se.cell(0, i).paragraphs[0].text = h
    format_row(tbl_se.rows[0], "F1F5F9", is_header=True)

    se_data = [
        ("1. Communication", "Identify stakeholder needs and resolve domain ambiguities.", "Defined retail pricing needs; established soft-archival rule to preserve history; handled zero-price promotions.", "SRS document, feature matrix, acceptance criteria."),
        ("2. Planning", "Map project tasks, dependencies, and risk mitigation.", "Prioritized core catalog & pricing before advanced analytics; mitigated CSV write corruption with atomic moves.", "Milestone schedule, risk matrix, task allocation."),
        ("3. Modelling", "Translate requirements into architectural & data designs.", "Engineered 5-layer design, Product/PriceHistory domain models, RFC-4180 CSV schemas, and state flows.", "Class diagrams, use-case models, CSV schemas."),
        ("4. Construction", "Transform approved models into verified Java code.", "Implemented encapsulated classes, BigDecimal math, CsvFileHandler atomic writes, and 25 automated unit tests.", "Compilable Java source code, TestRunner suite."),
        ("5. Deployment", "Package, verify, and document for demonstration.", "Created cross-platform scripts (compile.bat, run.bat, build.ps1); automated demo data seeding in Application.java.", "Executable binaries, user guide, final project report.")
    ]
    for r_idx, (act, obj, imp, deliv) in enumerate(se_data):
        row = tbl_se.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = act
        row.cells[1].paragraphs[0].text = obj
        row.cells[2].paragraphs[0].text = imp
        row.cells[3].paragraphs[0].text = deliv
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # 6. USERS AND RESPONSIBILITIES
    # =========================================================================
    add_header(doc, "6. Users and Responsibilities")
    
    tbl_users = doc.add_table(rows=5, cols=2)
    tbl_users.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_users.autofit = False
    tbl_users.columns[0].width = Inches(2.2)
    tbl_users.columns[1].width = Inches(4.5)
    set_table_borders(tbl_users, color="CBD5E1", sz="4")

    u_headers = ["User Role", "Main Functions & Responsibilities"]
    for i, h in enumerate(u_headers):
        tbl_users.cell(0, i).paragraphs[0].text = h
    format_row(tbl_users.rows[0], "F1F5F9", is_header=True)

    users_data = [
        ("Store Administrator", "Oversee catalog health, configure system settings, trigger backups/restores, and generate executive valuation reports."),
        ("Pricing Manager", "Authorize selling price adjustments, review proposed percentage changes, provide audit justifications, and execute bulk revaluations."),
        ("Inventory Clerk", "Register new merchandise, update cost prices, record physical stock quantities, and monitor low-stock restock alerts."),
        ("Financial Auditor", "Inspect chronological price audit trails, examine margin reports, identify loss-leader items, and export audit data to CSV.")
    ]
    for r_idx, (u, f) in enumerate(users_data):
        row = tbl_users.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = u
        row.cells[1].paragraphs[0].text = f
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # 7. SUGGESTED TECHNOLOGY STACK
    # =========================================================================
    add_header(doc, "7. Suggested Technology Stack")
    
    tbl_tech = doc.add_table(rows=8, cols=2)
    tbl_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_tech.autofit = False
    tbl_tech.columns[0].width = Inches(2.2)
    tbl_tech.columns[1].width = Inches(4.5)
    set_table_borders(tbl_tech, color="CBD5E1", sz="4")

    t_headers = ["System Layer / Component", "Technology Choice & Justification"]
    for i, h in enumerate(t_headers):
        tbl_tech.cell(0, i).paragraphs[0].text = h
    format_row(tbl_tech.rows[0], "F1F5F9", is_header=True)

    tech_data = [
        ("Programming Language", "Core Java (OpenJDK 17 LTS) — Standard library only"),
        ("Architecture Pattern", "Decoupled 5-Layer Modular Object-Oriented Architecture"),
        ("Persistence Engine", "Local File Storage (RFC-4180 CSV files via java.nio.file)"),
        ("Financial Math Engine", "java.math.BigDecimal with explicit RoundingMode.HALF_UP"),
        ("User Interface", "Advanced Interactive Console UI with ASCII Borders & ANSI Colors"),
        ("Testing Framework", "Self-Contained Automated Test Suite (Assert.java & TestRunner.java)"),
        ("Build & Automation", "Native shell scripts (compile.bat, run.bat, build.ps1, compile.sh)")
    ]
    for r_idx, (l, t) in enumerate(tech_data):
        row = tbl_tech.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = l
        row.cells[1].paragraphs[0].text = t
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # 8. IMPORTANT DATA STORAGE SCHEMAS
    # =========================================================================
    add_header(doc, "8. Important Data Storage Schemas")
    add_p(doc, "PPMS maintains persistent data across three normalized local CSV files in the data/ directory:")

    add_box_diagram(doc,
"PRODUCTS (data/products.csv)\n"
"-------------------------------------------------------------------------\n"
"productId      : Alphanumeric permanent ID (PRD-xxxx)\n"
"name           : Product commercial name (quoted if containing commas)\n"
"category       : Product classification / department\n"
"description    : Product technical/marketing description\n"
"costPrice      : Normalized procurement cost (BigDecimal, 2 decimals)\n"
"sellingPrice   : Active retail selling price (BigDecimal, 2 decimals)\n"
"quantity       : Units currently in physical inventory (int >= 0)\n"
"status         : Lifecycle state (ACTIVE / ARCHIVED)\n"
"createdAt      : ISO-8601 registration timestamp (yyyy-MM-dd'T'HH:mm:ss)\n"
"updatedAt      : ISO-8601 modification timestamp (yyyy-MM-dd'T'HH:mm:ss)\n\n"
"PRICE_HISTORY (data/price_history.csv)\n"
"-------------------------------------------------------------------------\n"
"historyId      : Permanent unique audit record ID (HST-xxxx)\n"
"productId      : Foreign reference to product (PRD-xxxx)\n"
"oldPrice       : Selling price immediately prior to modification\n"
"newPrice       : Approved new selling price\n"
"changedAt      : ISO-8601 price alteration timestamp\n"
"reason         : Mandatory business reason for price revision\n\n"
"METADATA (data/metadata.csv)\n"
"-------------------------------------------------------------------------\n"
"key            : Configuration attribute name (e.g., low_stock_threshold)\n"
"value          : Persisted operational value (e.g., 10, Rs., 1.0.0)"
    )

    # =========================================================================
    # 9. MAIN WORKFLOW & MAJOR USE CASES
    # =========================================================================
    add_header(doc, "9. Main Workflow & Major Use Cases")
    add_p(doc, "The operational lifecycle follows a structured sequence:", bold_prefix="Main Pricing & Operational Workflow: ")

    add_box_diagram(doc,
"Application Launch & Verification\n"
"          |\n"
"          v\n"
"High-Watermark Sequence Counter Initialization\n"
"          |\n"
"          v\n"
"Load Products & Price History from CSV Cache\n"
"          |\n"
"          v\n"
"Calculate Real-Time Dashboard Statistics & Anomaly Alerts\n"
"          |\n"
"          v\n"
"Select Menu Action (Catalog / Pricing / History / Analytics / Backup)\n"
"          |\n"
"          v\n"
"Input Validation & Domain Boundary Verification\n"
"          |\n"
"          v\n"
"Generate Proposal Preview (Display Diff, % Change, Margin Warnings)\n"
"          |\n"
"          v\n"
"Explicit User Confirmation Gateway (Yes / No)\n"
"          |\n"
"          v\n"
"Update Product Entity & Create Immutable PriceHistory Audit Record\n"
"          |\n"
"          v\n"
"Atomic Safe Write to Local CSV (.tmp staging -> .bak -> atomic rename)\n"
"          |\n"
"          v\n"
"Display Success Notice & Return to Interactive Menu"
    )

    add_p(doc, "Table 9.1 documents the major use cases supported by the system:")
    tbl_mu = doc.add_table(rows=5, cols=2)
    tbl_mu.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mu.autofit = False
    tbl_mu.columns[0].width = Inches(2.2)
    tbl_mu.columns[1].width = Inches(4.5)
    set_table_borders(tbl_mu, color="CBD5E1", sz="4")

    mu_headers = ["Actor", "Supported System Use Cases"]
    for i, h in enumerate(mu_headers):
        tbl_mu.cell(0, i).paragraphs[0].text = h
    format_row(tbl_mu.rows[0], "F1F5F9", is_header=True)

    mu_data = [
        ("Store Administrator", "View Dashboard, Generate Executive Valuation Report, Configure System Settings, Create Snapshot Backup, Restore Catalog from Backup."),
        ("Pricing Manager", "Inspect Current Price, Adjust Selling Price, Preview % Change, Submit Audit Reason, Confirm Price Update, Execute Bulk Category Adjustment."),
        ("Inventory Clerk", "Add Product, Inspect Product Details, Search Catalog, Filter Products, Sort Products, Update Cost Price, Update Quantity, Soft-Archive Product."),
        ("Financial Auditor", "View Full Price Audit Log, Filter History by Product ID, Filter History by Date Range, Inspect Largest Increase/Decrease, Export CSV Data.")
    ]
    for r_idx, (act, ucs) in enumerate(mu_data):
        row = tbl_mu.rows[r_idx + 1]
        row.cells[0].paragraphs[0].text = act
        row.cells[1].paragraphs[0].text = ucs
        format_row(row, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # =========================================================================
    # 10. FUNCTIONAL REQUIREMENTS
    # =========================================================================
    add_header(doc, "10. Functional Requirements")
    add_bullet(doc, "Real-time calculated statistics dashboard with zero hardcoding")
    add_bullet(doc, "Product creation with automated collision-free ID generation (PRD-xxxx)")
    add_bullet(doc, "Multi-attribute search across ID, name, category, and description")
    add_bullet(doc, "Multi-criteria filtering with price and quantity boundary validation")
    add_bullet(doc, "Deterministic multi-field sorting (Name, Category, Price, Cost, Stock)")
    add_bullet(doc, "Soft-archival and restoration protecting historical pricing integrity")
    add_bullet(doc, "Selling price modification preview with signed difference and % change")
    add_bullet(doc, "Mandatory business reason prompt before approving price modifications")
    add_bullet(doc, "Immutable price change history logging (historyId, oldPrice, newPrice, reason)")
    add_bullet(doc, "Chronological audit querying with date-range and keyword filters")
    add_bullet(doc, "Real-time calculation of median price, inventory cost, and retail valuation")
    add_bullet(doc, "Automated alerts for loss-leader products sold below cost price")
    add_bullet(doc, "Automated warning identification for low-stock and depleted inventory")
    add_bullet(doc, "Bulk percentage price adjustments with tabular pre-commit preview")
    add_bullet(doc, "RFC-4180 compliant CSV export for products and price history")
    add_bullet(doc, "Batch CSV product import with line-by-line validation and rejection logs")
    add_bullet(doc, "Timestamped snapshot backup creation and safe pre-validated restoration")

    # =========================================================================
    # 11. NON-FUNCTIONAL REQUIREMENTS
    # =========================================================================
    add_header(doc, "11. Non-Functional Requirements")
    add_bullet(doc, "Financial Accuracy: Exclusive use of java.math.BigDecimal with RoundingMode.HALF_UP; zero floating-point errors.")
    add_bullet(doc, "Data Integrity: Crash-safe atomic file replacement (.tmp staging -> .bak backup -> atomic rename).")
    add_bullet(doc, "Reliability: High-watermark atomic sequence generation prevents duplicate or recycled IDs across restarts.")
    add_bullet(doc, "Usability: Clean, aligned ASCII table borders (+---, |), paginated viewing, and toggleable ANSI colors.")
    add_bullet(doc, "Self-Sufficiency: 100% offline standalone execution requiring zero network connectivity or database servers.")
    add_bullet(doc, "Maintainability: Decoupled 5-layer design with interface-driven repository abstractions.")
    add_bullet(doc, "Portability: Standard Core Java SE 17 runs identically on Windows, Linux, and macOS.")

    # =========================================================================
    # 12. SOFTWARE ENGINEERING DOCUMENTATION & TEST VERIFICATION
    # =========================================================================
    add_header(doc, "12. Software Engineering Documentation & Test Verification")
    add_p(doc, "As part of complete Software Engineering academic delivery, the project includes:")
    add_bullet(doc, "Problem Definition & Feasibility Study")
    add_bullet(doc, "Software Requirements Specification (SRS) with Functional and Non-Functional Requirements")
    add_bullet(doc, "Use-Case Model & Actor Specifications")
    add_bullet(doc, "Class Diagram & Domain Model Design")
    add_bullet(doc, "Data Flow & Master Lifecycle Flowcharts")
    add_bullet(doc, "CSV Schemas & Data Dictionary (products.csv, price_history.csv, metadata.csv)")
    add_bullet(doc, "Console UI Design & Tabular Mockups")
    add_bullet(doc, "System Layered Architecture Specification")
    add_bullet(doc, "Test Plan & 25 Automated Unit/Integration Test Cases (100% Pass Rate Verified)")
    add_bullet(doc, "Cross-Platform Build & Execution Scripts (compile.bat, run.bat, build.ps1)")
    add_bullet(doc, "Final Comprehensive Software Engineering Report")

    # =========================================================================
    # 13. EXPECTED OUTCOME
    # =========================================================================
    add_header(doc, "13. Expected Outcome")
    add_p(doc, "The completed system provides a standalone, robust Core Java desktop platform through which store managers and inventory administrators can manage product catalogs, enforce disciplined pricing governance, record every relevant price modification in an immutable audit ledger, examine historical pricing trends, calculate gross margins and inventory valuations in real time, and persist all records locally using crash-safe CSV files. The application eliminates reliance on complex external database engines and ensures complete operational reliability and financial accuracy offline.")

    # Save document with dedicated filename
    output_path = os.path.abspath("Product_Price_Management_System_Project_Report_Formatted.docx")
    doc.save(output_path)
    print(f"[SUCCESS] Exact layout report generated at: {output_path}")

    # Also save as Product_Price_Management_System_Report.docx for convenience
    try:
        alt_path = os.path.abspath("Product_Price_Management_System_Report.docx")
        doc.save(alt_path)
        print(f"[SUCCESS] Also saved at: {alt_path}")
    except Exception as e:
        print(f"Notice: {e}")

if __name__ == "__main__":
    generate_report()
